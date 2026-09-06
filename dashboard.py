"""
Sinyalci — Streamlit Web Dashboard.

Modern koyu temalı arayüz: coin takibi, canlı fiyat, sinyal ve risk seviyeleri.

Çalıştırma:
    python -m streamlit run dashboard.py

Not (Windows): streamlit komutu PATH'te değilse yukarıdaki komutu kullanın.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime, timezone
from html import escape
from io import BytesIO
import time
from textwrap import dedent
from typing import Any, Optional

import streamlit as st
import streamlit.components.v1 as components
import pandas as pd

from src import (
    ScanOpportunity,
    ScannerConfig,
    SignalType,
    build_risk_payload_from_signal,
    evaluate_signal,
    fetch_ohlcv,
    fetch_live_tickers,
    fetch_price_change_pct,
    opportunities_to_dataframe,
    scan_market,
)
from src.market_scanner import ScanStats, is_blacklisted_symbol
from src.risk_manager import calculate_exit_levels, estimate_target_eta
from src.recommendation_store import (
    clear_recommendation_history,
    get_history_file_path,
    load_recommendation_history,
    save_recommendation_history,
)
from src.watchlist_store import (
    load_watchlist_positions,
    upsert_watchlist_position,
)

# Kart alanı otomatik yenileme — yalnızca @st.fragment, tam sayfa rerun yok
REFRESH_INTERVAL_SEC = 10

# Coin başına kısa bekleme — API rate limit koruması
API_REQUEST_DELAY_SEC = 0.15

# Piyasa radarında saklanacak son tavsiye sayısı
RECOMMENDATION_HISTORY_LIMIT = 5
REASON_TECHNICAL_MARKER = "Koşul özeti:"


# ---------------------------------------------------------------------------
# Yardımcı veri yapıları
# ---------------------------------------------------------------------------


@dataclass
class CoinSnapshot:
    """Tek bir coin için dashboard verisi."""

    symbol: str
    live_price: float
    signal_label: str
    signal_type: SignalType
    rsi: Optional[float]
    risk_payload: Optional[dict[str, Any]]
    updated_at: Optional[datetime] = None
    is_stale: bool = False
    error: Optional[str] = None


# ---------------------------------------------------------------------------
# Özel CSS — koyu tema, gölgeli ve yuvarlatılmış kartlar
# ---------------------------------------------------------------------------


def render_html(html: str, height: Optional[int] = None) -> None:
    """
    HTML içeriğini güvenli şekilde render eder.

    Streamlit markdown'da ``$`` işareti LaTeX olarak yorumlanabildiği ve
    girintili HTML kod bloğu sayılabildiği için kartlar ``components.html``
    ile basılır; diğer küçük bloklar ``st.markdown`` kullanır.
    """
    cleaned = dedent(html).strip()
    if height is not None:
        components.html(cleaned, height=height, scrolling=False)
    else:
        st.markdown(cleaned, unsafe_allow_html=True)


def hide_streamlit_chrome() -> None:
    """Deploy butonu, üç nokta menüsü ve varsayılan üst/alt çubuğu gizler."""
    st.markdown(
        "<style>#MainMenu {visibility: hidden;} header {visibility: hidden;} "
        "footer {visibility: hidden;} .stDeployButton {display: none;}</style>",
        unsafe_allow_html=True,
    )


def inject_custom_css() -> None:
    """Streamlit arayüzüne modern koyu tema stillerini enjekte eder."""
    hide_streamlit_chrome()
    render_html(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

        /* Ana arka plan */
        .stApp {
            background: linear-gradient(145deg, #0a0a0f 0%, #12121a 40%, #1a1a2e 100%);
            font-family: 'Inter', sans-serif;
        }

        /* Streamlit varsayılan üst boşlukları */
        .block-container {
            padding-top: 2rem;
            max-width: 1200px;
        }

        /* Sidebar */
        section[data-testid="stSidebar"] {
            background: linear-gradient(180deg, #0d0d14 0%, #141420 100%);
            border-right: 1px solid rgba(255, 255, 255, 0.06);
        }
        section[data-testid="stSidebar"] .stMarkdown h1,
        section[data-testid="stSidebar"] .stMarkdown h2,
        section[data-testid="stSidebar"] .stMarkdown h3 {
            color: #e2e8f0;
        }

        /* Başlık alanı */
        .dashboard-header {
            text-align: center;
            padding: 1.5rem 0 2.5rem 0;
        }
        .dashboard-header h1 {
            font-size: 2.4rem;
            font-weight: 700;
            background: linear-gradient(90deg, #60a5fa, #a78bfa, #34d399);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            margin-bottom: 0.4rem;
        }
        .dashboard-header p {
            color: #94a3b8;
            font-size: 1rem;
            margin: 0;
        }

        /* Metrik kartı */
        div[data-testid="stMetric"] {
            background: rgba(255, 255, 255, 0.02);
            border-radius: 12px;
            padding: 0.35rem 0.25rem;
        }
        div[data-testid="stMetricValue"] {
            font-size: 0.95rem !important;
            line-height: 1.25 !important;
            overflow-wrap: anywhere;
        }
        div[data-testid="stMetricLabel"] {
            font-size: 0.72rem !important;
        }

        .metric-card {
            background: rgba(255, 255, 255, 0.03);
            backdrop-filter: blur(16px);
            -webkit-backdrop-filter: blur(16px);
            border: 1px solid rgba(255, 255, 255, 0.08);
            border-radius: 20px;
            padding: 1.75rem;
            margin-bottom: 1rem;
            box-shadow:
                0 4px 24px rgba(0, 0, 0, 0.35),
                0 1px 0 rgba(255, 255, 255, 0.04) inset;
            transition: transform 0.25s ease, box-shadow 0.25s ease;
        }
        .metric-card:hover {
            transform: translateY(-3px);
            box-shadow:
                0 12px 40px rgba(0, 0, 0, 0.45),
                0 1px 0 rgba(255, 255, 255, 0.06) inset;
        }

        .card-symbol {
            font-size: 1.25rem;
            font-weight: 700;
            color: #f1f5f9;
            letter-spacing: 0.02em;
        }
        .card-price {
            font-size: 2rem;
            font-weight: 700;
            color: #ffffff;
            margin: 0.6rem 0;
            font-variant-numeric: tabular-nums;
        }
        .card-price-label {
            font-size: 0.75rem;
            color: #64748b;
            text-transform: uppercase;
            letter-spacing: 0.08em;
        }

        /* Sinyal rozetleri */
        .signal-badge {
            display: inline-block;
            padding: 0.45rem 1rem;
            border-radius: 999px;
            font-size: 0.85rem;
            font-weight: 600;
            letter-spacing: 0.03em;
        }
        .signal-buy {
            background: rgba(16, 185, 129, 0.15);
            color: #34d399;
            border: 1px solid rgba(52, 211, 153, 0.35);
            box-shadow: 0 0 20px rgba(52, 211, 153, 0.15);
        }
        .signal-sell {
            background: rgba(239, 68, 68, 0.15);
            color: #f87171;
            border: 1px solid rgba(248, 113, 113, 0.35);
            box-shadow: 0 0 20px rgba(248, 113, 113, 0.15);
        }
        .signal-hold {
            background: rgba(148, 163, 184, 0.12);
            color: #cbd5e1;
            border: 1px solid rgba(148, 163, 184, 0.25);
        }

        .card-meta {
            color: #64748b;
            font-size: 0.8rem;
            margin-top: 0.75rem;
        }

        /* Çıkış stratejisi bölümü */
        .risk-section {
            margin-top: 1rem;
            padding-top: 1rem;
            border-top: 1px solid rgba(255, 255, 255, 0.06);
        }
        .risk-title {
            font-size: 0.7rem;
            font-weight: 600;
            color: #64748b;
            text-transform: uppercase;
            letter-spacing: 0.1em;
            margin-bottom: 0.75rem;
        }

        .risk-row {
            display: flex;
            flex-wrap: wrap;
            gap: 0.6rem;
        }

        .risk-chip {
            flex: 1;
            min-width: 140px;
            padding: 0.75rem 1rem;
            border-radius: 14px;
            font-size: 0.82rem;
        }
        .risk-chip strong {
            display: block;
            font-size: 1rem;
            margin-top: 0.2rem;
            font-variant-numeric: tabular-nums;
        }
        .risk-chip span {
            opacity: 0.85;
            font-size: 0.72rem;
        }

        .chip-profit {
            background: rgba(16, 185, 129, 0.12);
            border: 1px solid rgba(52, 211, 153, 0.3);
            color: #6ee7b7;
        }
        .chip-stop {
            background: rgba(239, 68, 68, 0.12);
            border: 1px solid rgba(248, 113, 113, 0.3);
            color: #fca5a5;
        }
        .eta-note {
            display: inline-block;
            margin-top: 0.4rem;
            padding: 0.15rem 0.5rem;
            border-radius: 999px;
            background: rgba(255, 255, 255, 0.08);
            font-size: 0.68rem;
            font-weight: 600;
            letter-spacing: 0.01em;
        }

        /* Hata kartı */
        .error-card {
            background: rgba(239, 68, 68, 0.08);
            border: 1px solid rgba(248, 113, 113, 0.25);
            border-radius: 16px;
            padding: 1.25rem;
            color: #fca5a5;
        }

        /* Durum çubuğu */
        .status-bar {
            display: flex;
            justify-content: center;
            align-items: center;
            gap: 0.5rem;
            padding: 0.55rem 1.2rem;
            margin-bottom: 1.75rem;
            background: rgba(255, 255, 255, 0.04);
            border: 1px solid rgba(255, 255, 255, 0.07);
            border-radius: 999px;
            color: #94a3b8;
            font-size: 0.82rem;
            width: fit-content;
            margin-left: auto;
            margin-right: auto;
        }
        .status-dot {
            width: 8px;
            height: 8px;
            border-radius: 50%;
            background: #34d399;
            box-shadow: 0 0 10px rgba(52, 211, 153, 0.6);
            animation: pulse 2s infinite;
        }
        .status-dot.stale {
            background: #fbbf24;
            box-shadow: 0 0 10px rgba(251, 191, 36, 0.5);
        }
        @keyframes pulse {
            0%, 100% { opacity: 1; }
            50% { opacity: 0.45; }
        }

        .stale-banner {
            background: rgba(251, 191, 36, 0.1);
            border: 1px solid rgba(251, 191, 36, 0.28);
            border-radius: 10px;
            padding: 0.45rem 0.75rem;
            color: #fcd34d;
            font-size: 0.75rem;
            margin-bottom: 0.75rem;
        }

        /* Streamlit arayüz öğelerini gizle */
        #MainMenu { visibility: hidden; }
        header { visibility: hidden; }
        footer { visibility: hidden; }
        .stDeployButton { display: none; }

        /* Footer gizleme — yedek */
        footer { visibility: hidden; }
        header[data-testid="stHeader"] {
            background: transparent;
        }
        </style>
        """
    )


# ---------------------------------------------------------------------------
# Veri çekme ve görselleştirme
# ---------------------------------------------------------------------------


def parse_symbols(raw: str) -> list[str]:
    """Virgülle ayrılmış coin listesini temizler ve doğrular."""
    symbols = []
    for part in raw.split(","):
        symbol = part.strip().upper()
        if symbol and symbol not in symbols:
            symbols.append(symbol)
    return symbols


def format_price(price: float) -> str:
    """Fiyat büyüklüğüne göre okunabilir format üretir (USDT)."""
    if price <= 0:
        return "—"
    if price >= 1000:
        return f"{price:,.2f} USDT"
    if price >= 1:
        return f"{price:,.4f} USDT"
    return f"{price:,.6f} USDT"


def card_styles() -> str:
    """Kart bileşenleri için iframe içi CSS (components.html ile kullanılır)."""
    return """
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
        html, body {
            margin: 0;
            padding: 0;
            background: transparent;
            font-family: 'Inter', sans-serif;
        }
        .metric-card {
            background: rgba(255, 255, 255, 0.03);
            border: 1px solid rgba(255, 255, 255, 0.08);
            border-radius: 20px;
            padding: 1.75rem;
            box-shadow: 0 4px 24px rgba(0, 0, 0, 0.35);
        }
        .card-symbol { font-size: 1.25rem; font-weight: 700; color: #f1f5f9; }
        .card-price-label {
            font-size: 0.75rem; color: #64748b;
            text-transform: uppercase; letter-spacing: 0.08em; margin-top: 0.5rem;
        }
        .card-price {
            font-size: 2rem; font-weight: 700; color: #ffffff;
            margin: 0.6rem 0; font-variant-numeric: tabular-nums;
        }
        .signal-badge {
            display: inline-block; padding: 0.45rem 1rem; border-radius: 999px;
            font-size: 0.85rem; font-weight: 600;
        }
        .signal-buy {
            background: rgba(16, 185, 129, 0.15); color: #34d399;
            border: 1px solid rgba(52, 211, 153, 0.35);
        }
        .signal-sell {
            background: rgba(239, 68, 68, 0.15); color: #f87171;
            border: 1px solid rgba(248, 113, 113, 0.35);
        }
        .signal-hold {
            background: rgba(148, 163, 184, 0.12); color: #cbd5e1;
            border: 1px solid rgba(148, 163, 184, 0.25);
        }
        .card-meta { color: #64748b; font-size: 0.8rem; margin-top: 0.75rem; }
        .stale-banner {
            background: rgba(251, 191, 36, 0.1);
            border: 1px solid rgba(251, 191, 36, 0.28);
            border-radius: 10px; padding: 0.45rem 0.75rem;
            color: #fcd34d; font-size: 0.75rem; margin-bottom: 0.75rem;
        }
        .error-card {
            background: rgba(239, 68, 68, 0.08);
            border: 1px solid rgba(248, 113, 113, 0.25);
            border-radius: 16px; padding: 1.25rem; color: #fca5a5;
        }
        .risk-section {
            margin-top: 1rem; padding-top: 1rem;
            border-top: 1px solid rgba(255, 255, 255, 0.06);
        }
        .risk-title {
            font-size: 0.7rem; font-weight: 600; color: #64748b;
            text-transform: uppercase; letter-spacing: 0.1em; margin-bottom: 0.75rem;
        }
        .risk-row { display: flex; flex-wrap: wrap; gap: 0.6rem; }
        .risk-chip {
            flex: 1; min-width: 120px; padding: 0.75rem 1rem;
            border-radius: 14px; font-size: 0.82rem;
        }
        .risk-chip strong {
            display: block; font-size: 1rem; margin-top: 0.2rem;
            font-variant-numeric: tabular-nums;
        }
        .risk-chip span { opacity: 0.85; font-size: 0.72rem; }
        .chip-profit {
            background: rgba(16, 185, 129, 0.12);
            border: 1px solid rgba(52, 211, 153, 0.3); color: #6ee7b7;
        }
        .chip-stop {
            background: rgba(239, 68, 68, 0.12);
            border: 1px solid rgba(248, 113, 113, 0.3); color: #fca5a5;
        }
        .eta-note {
            display: inline-block; margin-top: 0.4rem; padding: 0.15rem 0.5rem;
            border-radius: 999px; background: rgba(255, 255, 255, 0.08);
            font-size: 0.68rem; font-weight: 600;
        }
        .opportunity-card {
            background: linear-gradient(135deg, rgba(16,185,129,0.08) 0%, rgba(255,255,255,0.03) 100%);
            border: 1px solid rgba(52, 211, 153, 0.35);
            border-radius: 20px;
            padding: 1.5rem;
            box-shadow: 0 8px 32px rgba(16, 185, 129, 0.12);
            margin-bottom: 1rem;
        }
        .opp-header {
            display: flex; justify-content: space-between; align-items: center;
            margin-bottom: 1rem;
        }
        .opp-symbol {
            font-size: 1.35rem; font-weight: 700; color: #6ee7b7;
        }
        .opp-badge {
            background: rgba(16,185,129,0.2); color: #34d399;
            padding: 0.35rem 0.85rem; border-radius: 999px;
            font-size: 0.78rem; font-weight: 600;
        }
        .opp-grid {
            display: grid; grid-template-columns: repeat(2, 1fr); gap: 0.75rem;
            margin-bottom: 1rem;
        }
        .opp-stat {
            background: rgba(0,0,0,0.2); border-radius: 12px; padding: 0.75rem;
        }
        .opp-stat label {
            display: block; font-size: 0.68rem; color: #64748b;
            text-transform: uppercase; letter-spacing: 0.06em;
        }
        .opp-stat value {
            display: block; font-size: 0.95rem; font-weight: 600;
            color: #f1f5f9; margin-top: 0.25rem;
            font-variant-numeric: tabular-nums;
        }
        .opp-stat value.green { color: #6ee7b7; }
        .opp-stat value.red { color: #fca5a5; }
        .opp-reason {
            background: rgba(16,185,129,0.06); border-left: 3px solid #34d399;
            padding: 0.75rem 1rem; border-radius: 0 10px 10px 0;
            color: #94a3b8; font-size: 0.82rem; line-height: 1.5;
        }
    </style>
    """


def utc_now() -> datetime:
    """Timezone-aware UTC zaman damgası döndürür."""
    return datetime.now(timezone.utc)


def format_datetime_tr(dt: datetime) -> str:
    """Zaman damgasını Türkçe arayüz için biçimlendirir."""
    return dt.strftime("%d.%m.%Y %H:%M:%S")


def get_signal_badge_class(signal_type: SignalType) -> str:
    """Sinyal tipine göre CSS sınıfı döndürür."""
    if signal_type in (SignalType.LONG, SignalType.BUY):
        return "signal-buy"
    if signal_type in (SignalType.SHORT, SignalType.SELL):
        return "signal-sell"
    return "signal-hold"


def get_signal_short_label(signal_type: SignalType) -> str:
    """Arayüzde gösterilecek kısa sinyal etiketi."""
    if signal_type in (SignalType.LONG, SignalType.BUY):
        return "LONG"
    if signal_type in (SignalType.SHORT, SignalType.SELL):
        return "SHORT"
    return "BEKLE"


def fetch_coin_snapshot(
    symbol: str,
    interval: str,
    history_limit: int,
) -> CoinSnapshot:
    """
    Tek coin için data_engine + strategy_engine + risk_manager verisini çeker.

    Her adım ayrı try-except ile korunur; kısmi hatalarda mümkün olduğunca
    veri döndürülür.
    """
    now = utc_now()

    try:
        # Geçmiş OHLCV verisi — son kapanış fiyatı canlı fiyat olarak kullanılır
        df = fetch_ohlcv(symbol=symbol, interval=interval, limit=history_limit)
        live_price = float(df["Close"].iloc[-1])
    except Exception as exc:
        return CoinSnapshot(
            symbol=symbol,
            live_price=0.0,
            signal_label="—",
            signal_type=SignalType.BEKLE,
            rsi=None,
            risk_payload=None,
            updated_at=now,
            error="Veri alınamadı. Bağlantınızı kontrol edip tekrar deneyin.",
        )

    try:
        signal_result = evaluate_signal(
            df,
            price_change_pct_24h=fetch_price_change_pct(symbol),
        )
    except Exception as exc:
        return CoinSnapshot(
            symbol=symbol,
            live_price=live_price,
            signal_label="—",
            signal_type=SignalType.BEKLE,
            rsi=None,
            risk_payload=None,
            updated_at=now,
            error="Sinyal hesaplanamadı. Veriler kısmen gösteriliyor.",
        )

    risk_payload: Optional[dict[str, Any]] = None
    try:
        risk_payload = build_risk_payload_from_signal(
            signal_result,
            symbol=symbol,
            interval=interval,
        )
    except Exception:
        # Risk seviyesi hesaplanamazsa kart yine de gösterilir
        risk_payload = None

    return CoinSnapshot(
        symbol=symbol,
        live_price=live_price,
        signal_label=signal_result.signal.value,
        signal_type=signal_result.signal,
        rsi=signal_result.rsi,
        risk_payload=risk_payload,
        updated_at=now,
    )


def refresh_all_snapshots(
    symbols: list[str],
    interval: str,
    history_limit: int,
) -> dict[str, CoinSnapshot]:
    """
    Coin listesini döngüyle gezer; her coin için güncel snapshot üretir.

    API rate limit riskini azaltmak için istekler arasında kısa gecikme uygular.
    """
    snapshots: dict[str, CoinSnapshot] = {}
    previous = st.session_state.get("snapshots", {})

    for index, symbol in enumerate(symbols):
        try:
            snapshot = fetch_coin_snapshot(
                symbol=symbol,
                interval=interval,
                history_limit=history_limit,
            )
        except Exception as exc:
            snapshot = CoinSnapshot(
                symbol=symbol,
                live_price=0.0,
                signal_label="—",
                signal_type=SignalType.BEKLE,
                rsi=None,
                risk_payload=None,
                updated_at=utc_now(),
                error="Beklenmeyen bir sorun oluştu.",
            )

        # Hata durumunda önceki başarılı veriyi koru (UI çökmesin)
        if snapshot.error and symbol in previous and previous[symbol].error is None:
            snapshots[symbol] = replace(
                previous[symbol],
                is_stale=True,
                error=f"{snapshot.error} · Önceki veri gösteriliyor.",
            )
        else:
            snapshots[symbol] = snapshot

        # Son coin hariç istekler arası kısa bekleme
        if index < len(symbols) - 1:
            time.sleep(API_REQUEST_DELAY_SEC)

    return snapshots


def split_reason_text(reason: str) -> tuple[str, str]:
    """Gerekçeyi okunabilir özet ve teknik detay olarak ayırır."""
    if not reason:
        return "", ""
    if REASON_TECHNICAL_MARKER in reason:
        summary, technical = reason.split(REASON_TECHNICAL_MARKER, 1)
        return summary.strip(), f"{REASON_TECHNICAL_MARKER}{technical.strip()}"
    return reason.strip(), ""


def render_reason_block(reason: str) -> None:
    """Gerekçe metnini tam gösterir; teknik detayları isteğe bağlı expander'da açar."""
    summary, technical = split_reason_text(reason)
    if summary:
        st.markdown(
            f"<p style='font-style: italic; color: #94a3b8; line-height: 1.65; "
            f"margin: 0.25rem 0; white-space: pre-wrap; word-break: break-word;'>"
            f"{escape(summary)}</p>",
            unsafe_allow_html=True,
        )
    if technical:
        with st.expander("Teknik koşul detayları"):
            st.markdown(
                f"<p style='color: #94a3b8; font-size: 0.82rem; line-height: 1.55; "
                f"white-space: pre-wrap; word-break: break-word;'>{escape(technical)}</p>",
                unsafe_allow_html=True,
            )


def ensure_watchlist_loaded() -> None:
    """Session state'e diskten izleme listesi pozisyonlarını yükler."""
    if st.session_state.get("watchlist_loaded"):
        return
    st.session_state.watchlist_positions = load_watchlist_positions()
    st.session_state.watchlist_loaded = True


def _add_symbol_to_watchlist(symbols: list[str], symbol: str) -> list[str]:
    """Pariteyi izleme listesine ekler (tekrarsız)."""
    symbol = symbol.upper()
    if symbol not in symbols:
        return symbols + [symbol]
    return symbols


def render_add_to_watchlist_form(
    symbol: str,
    default_entry: float,
    default_target: float,
    default_stop: Optional[float] = None,
    signal_label: str = "LONG",
    key_prefix: str = "item",
) -> None:
    """İzleme listesine ekleme butonu ve alış/hedef formunu gösterir."""
    symbol = symbol.upper()
    form_open = (
        st.session_state.get("watchlist_form_symbol") == symbol
        and st.session_state.get("watchlist_form_prefix") == key_prefix
    )

    if not form_open:
        if st.button(
            "👀 İzleme Listesine Ekle",
            key=f"watchlist_add_btn_{key_prefix}_{symbol}",
            use_container_width=True,
        ):
            st.session_state.watchlist_form_symbol = symbol
            st.session_state.watchlist_form_prefix = key_prefix
            st.rerun()
        return

    safe_entry = default_entry if default_entry > 0 else 0.0001
    safe_target = default_target if default_target > 0 else safe_entry * 1.02

    with st.form(key=f"watchlist_form_{key_prefix}_{symbol}"):
        st.markdown(f"**{symbol}** — izleme listesine ekle")
        entry_price = st.number_input(
            "Alış Fiyatı (USDT)",
            min_value=0.0,
            value=float(safe_entry),
            format="%.8f",
            help="Pozisyona girdiğiniz veya girmeyi planladığınız fiyat.",
        )
        target_price = st.number_input(
            "Kar Hedefi (USDT)",
            min_value=0.0,
            value=float(safe_target),
            format="%.8f",
            help="Hedeflediğiniz çıkış fiyatı.",
        )
        stop_default = float(default_stop) if default_stop and default_stop > 0 else 0.0
        stop_loss = st.number_input(
            "Zarar Durdur (USDT) — isteğe bağlı",
            min_value=0.0,
            value=stop_default,
            format="%.8f",
        )

        col_ok, col_cancel = st.columns(2)
        with col_ok:
            submitted = st.form_submit_button("✅ Listeye Ekle", type="primary")
        with col_cancel:
            cancelled = st.form_submit_button("İptal")

    if cancelled:
        st.session_state.pop("watchlist_form_symbol", None)
        st.session_state.pop("watchlist_form_prefix", None)
        st.rerun()

    if submitted:
        if entry_price <= 0:
            st.error("Alış fiyatı 0'dan büyük olmalıdır.")
            return
        if target_price <= 0:
            st.error("Kar hedefi 0'dan büyük olmalıdır.")
            return

        positions = upsert_watchlist_position(
            symbol=symbol,
            entry_price=entry_price,
            target_price=target_price,
            stop_loss=stop_loss if stop_loss > 0 else None,
            signal_label=signal_label,
        )
        st.session_state.watchlist_positions = positions
        st.session_state.symbols = _add_symbol_to_watchlist(
            st.session_state.get("symbols", []),
            symbol,
        )
        st.session_state.pop("watchlist_form_symbol", None)
        st.session_state.pop("watchlist_form_prefix", None)
        st.session_state.pop("snapshots", None)
        st.success(
            f"**{symbol}** izleme listesine eklendi — "
            f"Alış: {format_price(entry_price)}, Hedef: {format_price(target_price)}"
        )
        st.rerun()


def ensure_recommendation_history_loaded() -> None:
    """Session state'e diskten tavsiye geçmişini bir kez yükler."""
    if st.session_state.get("recommendation_history_loaded"):
        return
    st.session_state.recommendation_history = load_recommendation_history(
        RECOMMENDATION_HISTORY_LIMIT
    )
    st.session_state.recommendation_history_loaded = True


def _opportunity_to_history_record(opp: ScanOpportunity) -> dict[str, Any]:
    """ScanOpportunity kaydını session geçmişi için sözlüğe çevirir."""
    raw = getattr(opp, "raw_signal", None)
    scanned = getattr(opp, "scanned_at", None) or utc_now()
    return {
        "symbol": opp.symbol,
        "signal_label": _opp_signal_label(opp),
        "tags": _format_opportunity_tags_text(opp),
        "entry_price": opp.entry_price,
        "current_price": opp.current_price,
        "stop_loss": opp.stop_loss,
        "take_profit_1": opp.take_profit_1,
        "take_profit_2": opp.take_profit_2,
        "reason": getattr(opp, "reason", None) or opp.recommendation or "",
        "price_change_pct_24h": getattr(opp, "price_change_pct_24h", None),
        "rsi": opp.rsi,
        "scanned_at": scanned.isoformat(),
        "is_early_long": bool(getattr(raw, "is_early_long", False)) if raw else False,
    }


def _update_recommendation_history(opportunities: list[ScanOpportunity]) -> None:
    """Yeni tavsiyeleri geçmişe ekler; son N kayıt tutulur (aynı parite güncellenir)."""
    if not opportunities:
        return

    history: list[dict[str, Any]] = list(
        st.session_state.get("recommendation_history", [])
    )

    for opp in reversed(opportunities):
        record = _opportunity_to_history_record(opp)
        history = [item for item in history if item.get("symbol") != opp.symbol]
        history.insert(0, record)

    trimmed = history[:RECOMMENDATION_HISTORY_LIMIT]
    st.session_state.recommendation_history = trimmed
    save_recommendation_history(trimmed, RECOMMENDATION_HISTORY_LIMIT)


def _parse_history_timestamp(raw: str) -> str:
    """ISO zaman damgasını arayüz formatına çevirir."""
    try:
        dt = datetime.fromisoformat(raw)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return format_datetime_tr(dt)
    except (TypeError, ValueError):
        return raw or "—"


def _render_compact_metric_row(metrics: list[tuple[str, str]]) -> None:
    """Dar alanlarda okunabilir, küçük fontlu metrik satırı basar."""
    cols = st.columns(len(metrics))
    for col, (label, value) in zip(cols, metrics):
        with col:
            st.markdown(
                f"<p style='margin:0;color:#64748b;font-size:0.72rem;"
                f"text-transform:uppercase;letter-spacing:0.04em;'>{escape(label)}</p>"
                f"<p style='margin:0.15rem 0 0 0;color:#f1f5f9;font-size:0.92rem;"
                f"font-weight:600;font-variant-numeric:tabular-nums;"
                f"word-break:break-word;line-height:1.3;'>{escape(value)}</p>",
                unsafe_allow_html=True,
            )


def render_recommendation_history() -> None:
    """Son kayıtlı tavsiyeleri işlem takibi için listeler."""
    ensure_recommendation_history_loaded()
    history: list[dict[str, Any]] = st.session_state.get("recommendation_history", [])

    header_left, header_right = st.columns([4, 1])
    with header_left:
        st.markdown("#### 📋 Son Tavsiyeler — İşlem Takibi")
    with header_right:
        if history and st.button(
            "Geçmişi Temizle",
            key="clear_recommendation_history",
            use_container_width=True,
        ):
            clear_recommendation_history()
            st.session_state.recommendation_history = []
            st.rerun()

    if not history:
        st.caption("Henüz kayıtlı tavsiye yok. İlk taramadan sonra son 5 sinyal buraya yazılır.")
        return

    st.caption(
        f"Son **{len(history)}** tavsiye kayıtlı (kalıcı dosya: "
        f"`{get_history_file_path().name}`). Yeni tarama yapsanız bile "
        f"bu liste korunur; tarayıcıyı kapatsanız bile diskten geri yüklenir."
    )

    for idx, item in enumerate(history):
        symbol = item.get("symbol", "—")
        signal_label = str(item.get("signal_label", "—")).upper()
        tags = item.get("tags", "")
        scanned_text = _parse_history_timestamp(str(item.get("scanned_at", "")))
        reason = item.get("reason", "")
        is_long = signal_label == "LONG"
        is_early = bool(item.get("is_early_long"))
        default_entry = float(item.get("entry_price") or item.get("current_price") or 0)
        default_target = float(item.get("take_profit_1") or default_entry)
        default_stop = item.get("stop_loss")
        default_stop_f = float(default_stop) if default_stop else None

        with st.container(border=True):
            header_cols = st.columns([2, 4, 2])
            with header_cols[0]:
                st.markdown(f"**{symbol}**")
                st.caption(f"📅 {scanned_text}")
            with header_cols[1]:
                if tags:
                    tag_parts = [part.strip() for part in tags.split("·") if part.strip()]
                    tag_cols = st.columns(min(len(tag_parts), 3) or 1)
                    for idx, tag in enumerate(tag_parts[:3]):
                        with tag_cols[idx]:
                            color = "green" if is_long else "red"
                            if "Aşırı Satış" in tag:
                                color = "orange"
                            if hasattr(st, "badge"):
                                st.badge(tag, color=color)
                            else:
                                st.markdown(tag)
            with header_cols[2]:
                if is_early:
                    st.warning("Erken LONG", icon="⚠️")
                elif is_long:
                    st.success("LONG", icon="✅")
                else:
                    st.error("SHORT", icon="🔻")

            entry_val = (
                format_price(float(item.get("entry_price") or 0))
                if item.get("entry_price")
                else "—"
            )
            sl_val = (
                format_price(float(item.get("stop_loss") or 0))
                if item.get("stop_loss")
                else "—"
            )
            tp1_val = (
                format_price(float(item.get("take_profit_1") or 0))
                if item.get("take_profit_1")
                else "—"
            )
            pct = item.get("price_change_pct_24h")
            pct_val = f"%{float(pct):+.2f}" if pct is not None else "—"

            _render_compact_metric_row(
                [
                    ("Giriş", entry_val),
                    ("Zarar Durdur", sl_val),
                    ("Kar Hedefi 1", tp1_val),
                    ("24s Değişim", pct_val),
                ]
            )

            if reason:
                render_reason_block(reason)

            render_add_to_watchlist_form(
                symbol=symbol,
                default_entry=default_entry,
                default_target=default_target,
                default_stop=default_stop_f,
                signal_label=signal_label,
                key_prefix=f"hist_{idx}",
            )

    st.markdown("---")


def render_risk_section(risk: dict[str, Any]) -> str:
    """LONG/SHORT sinyali için çıkış stratejisi HTML bloğu üretir."""
    tp1 = risk["take_profit_1"]
    tp2 = risk["take_profit_2"]
    sl = risk["stop_loss"]

    def _level_pct(level: dict[str, Any]) -> float:
        if "distance_pct" in level:
            return float(level["distance_pct"])
        if "percentage" in level:
            return float(level["percentage"])
        return 0.0

    tp1_pct = _level_pct(tp1)
    tp2_pct = _level_pct(tp2)
    sl_pct = _level_pct(sl)

    etiket_kar_1 = tp1.get("label") or "Kar Hedefi 1"
    etiket_kar_2 = tp2.get("label") or "Kar Hedefi 2"
    etiket_zarar = sl.get("label") or "Zarar Durdur"

    return f"""
<div class="risk-section">
    <div class="risk-title">Çıkış Stratejisi</div>
    <div class="risk-row">
        <div class="risk-chip chip-profit">
            <span>{escape(etiket_kar_1)} ({tp1_pct:+.1f}%)</span>
            <strong>{escape(format_price(tp1["price"]))}</strong>
            {f'<span class="eta-note">⏱ {escape(str(tp1.get("eta_text")))}</span>' if tp1.get("eta_text") else ""}
        </div>
        <div class="risk-chip chip-profit">
            <span>{escape(etiket_kar_2)} ({tp2_pct:+.1f}%)</span>
            <strong>{escape(format_price(tp2["price"]))}</strong>
            {f'<span class="eta-note">⏱ {escape(str(tp2.get("eta_text")))}</span>' if tp2.get("eta_text") else ""}
        </div>
        <div class="risk-chip chip-stop">
            <span>{escape(etiket_zarar)} ({sl_pct:+.1f}%)</span>
            <strong>{escape(format_price(sl["price"]))}</strong>
        </div>
    </div>
</div>
"""


def build_metric_card_html(snapshot: CoinSnapshot) -> str:
    """Metrik kartı HTML içeriğini oluşturur."""
    if snapshot.error and snapshot.live_price <= 0:
        return f"""
{card_styles()}
<div class="error-card">
    <strong>{escape(snapshot.symbol)}</strong><br>
    {escape(snapshot.error or "")}
</div>
"""

    badge_class = get_signal_badge_class(snapshot.signal_type)
    short_label = get_signal_short_label(snapshot.signal_type)
    rsi_text = f"{snapshot.rsi:.1f}" if snapshot.rsi is not None else "—"

    stale_html = ""
    if snapshot.error:
        stale_html = f'<div class="stale-banner">⚠ {escape(snapshot.error)}</div>'

    updated_text = (
        format_datetime_tr(snapshot.updated_at)
        if snapshot.updated_at
        else "—"
    )

    sinyal_metni = get_signal_short_label(snapshot.signal_type)

    risk_html = ""
    if snapshot.signal_type in (SignalType.LONG, SignalType.BUY, SignalType.SHORT, SignalType.SELL) and snapshot.risk_payload:
        risk_html = render_risk_section(snapshot.risk_payload)

    return f"""
{card_styles()}
<div class="metric-card">
    {stale_html}
    <div class="card-symbol">{escape(snapshot.symbol)}</div>
    <div class="card-price-label">Son Fiyat</div>
    <div class="card-price">{escape(format_price(snapshot.live_price))}</div>
    <span class="signal-badge {badge_class}">{escape(short_label)}</span>
    <div class="card-meta">
        RSI (14): {escape(rsi_text)} · Sinyal: {escape(sinyal_metni)} · Güncelleme: {escape(updated_text)}
    </div>
    {risk_html}
</div>
"""


def estimate_card_height(snapshot: CoinSnapshot) -> int:
    """Kart yüksekliğini içeriğe göre tahmin eder."""
    if snapshot.error and snapshot.live_price <= 0:
        return 120
    if snapshot.signal_type in (SignalType.LONG, SignalType.BUY, SignalType.SHORT, SignalType.SELL) and snapshot.risk_payload:
        return 430
    if snapshot.error:
        return 260
    return 220


def render_metric_card(snapshot: CoinSnapshot) -> None:
    """Tek coin için metrik kartını ekrana basar."""
    html = build_metric_card_html(snapshot)
    render_html(html, height=estimate_card_height(snapshot))


def opportunity_card_styles(is_long: bool) -> str:
    """Piyasa radarı fırsat kartları için iframe CSS (LONG/SHORT renk teması)."""
    if is_long:
        card_border = "rgba(52, 211, 153, 0.45)"
        card_bg = "linear-gradient(135deg, rgba(16,185,129,0.10) 0%, rgba(255,255,255,0.03) 100%)"
        badge_bg = "rgba(16,185,129,0.22)"
        badge_color = "#34d399"
        accent = "#6ee7b7"
        reason_border = "#34d399"
        reason_bg = "rgba(16,185,129,0.06)"
    else:
        card_border = "rgba(248, 113, 113, 0.45)"
        card_bg = "linear-gradient(135deg, rgba(239,68,68,0.10) 0%, rgba(255,255,255,0.03) 100%)"
        badge_bg = "rgba(239,68,68,0.22)"
        badge_color = "#f87171"
        accent = "#fca5a5"
        reason_border = "#f87171"
        reason_bg = "rgba(239,68,68,0.06)"

    return card_styles() + f"""
    <style>
        .opportunity-card {{
            background: {card_bg};
            border: 1px solid {card_border};
            border-radius: 20px;
            padding: 1.5rem;
            box-shadow: 0 8px 32px rgba(0, 0, 0, 0.25);
        }}
        .opp-header {{
            display: flex; justify-content: space-between; align-items: center;
            margin-bottom: 1rem;
        }}
        .opp-symbol {{ font-size: 1.35rem; font-weight: 700; color: {accent}; }}
        .opp-badge {{
            background: {badge_bg}; color: {badge_color};
            padding: 0.35rem 0.85rem; border-radius: 999px;
            font-size: 0.78rem; font-weight: 700;
        }}
        .opp-grid {{
            display: grid; grid-template-columns: repeat(2, 1fr); gap: 0.75rem;
            margin-bottom: 1rem;
        }}
        .opp-stat {{
            background: rgba(0,0,0,0.2); border-radius: 12px; padding: 0.75rem;
        }}
        .opp-stat label {{
            display: block; font-size: 0.68rem; color: #64748b;
            text-transform: uppercase; letter-spacing: 0.06em;
        }}
        .opp-stat span.val {{
            display: block; font-size: 0.95rem; font-weight: 600;
            color: #f1f5f9; margin-top: 0.25rem;
            font-variant-numeric: tabular-nums;
        }}
        .opp-stat span.val.green {{ color: #6ee7b7; }}
        .opp-stat span.val.red {{ color: #fca5a5; }}
        .opp-reason {{
            background: {reason_bg}; border-left: 3px solid {reason_border};
            padding: 0.75rem 1rem; border-radius: 0 10px 10px 0;
            color: #94a3b8; font-size: 0.82rem; line-height: 1.55;
        }}
        .opp-reason em {{
            font-style: italic; color: #cbd5e1; display: block; margin-top: 0.35rem;
        }}
        .eta-note {{
            display: inline-block; margin-top: 0.4rem; padding: 0.15rem 0.5rem;
            border-radius: 999px; background: rgba(255, 255, 255, 0.08);
            color: {accent}; font-size: 0.68rem; font-weight: 600;
        }}
    </style>
    """


def _opp_signal_label(opp: ScanOpportunity) -> str:
    """Fırsat kaydının sinyal etiketini döndürür."""
    return getattr(opp, "signal_label", None) or get_signal_short_label(opp.signal)


def _filter_radar_opportunities(
    opportunities: list[ScanOpportunity],
) -> list[ScanOpportunity]:
    """Stablecoin / kara liste paritelerini arayüzden tamamen çıkarır."""
    return [opp for opp in opportunities if not is_blacklisted_symbol(opp.symbol)]


def _get_opportunity_tags(opp: ScanOpportunity) -> list[tuple[str, str]]:
    """
    Fırsat için görsel etiket listesi döndürür.

    Returns
    -------
    list[tuple[str, str]]
        (etiket_metni, st.badge rengi) çiftleri.
    """
    signal_label = _opp_signal_label(opp)
    pct = float(getattr(opp, "price_change_pct_24h", 0) or 0)
    rsi = opp.rsi
    pool_sources = list(getattr(opp, "pool_sources", None) or [])
    raw = getattr(opp, "raw_signal", None)
    is_early_long = bool(getattr(raw, "is_early_long", False)) if raw else False

    tags: list[tuple[str, str]] = []

    if is_early_long or (
        signal_label == "LONG" and pct <= -10 and rsi is not None and rsi < 30
    ):
        tags.append(("🩸 Aşırı Satış Tepkisi", "orange"))

    if signal_label == "LONG" and (pct >= 10 or "trend" in pool_sources):
        if not any("Yükseliş" in label for label, _ in tags):
            tags.append(("🔥 Yükseliş Trendi", "green"))

    if signal_label == "SHORT" and (pct <= -10 or "loser" in pool_sources):
        tags.append(("📉 Düşüş Trendi", "red"))

    if signal_label == "SHORT" and pct >= 10:
        tags.append(("💨 Aşırı Alım Düzeltmesi", "red"))

    if "dinamik" in pool_sources:
        tags.append(("⚡ 24s dinamik hareket", "violet"))

    if "volume_surge" in pool_sources:
        tags.append(("📊 Hacim Patlaması", "blue"))

    if signal_label == "LONG" and not tags:
        tags.append(("🟢 Uzun Pozisyon", "green"))
    elif signal_label == "SHORT" and not tags:
        tags.append(("🔴 Kısa Pozisyon", "red"))

    # Yinelenen etiketleri kaldır, sırayı koru
    seen: set[str] = set()
    unique_tags: list[tuple[str, str]] = []
    for label, color in tags:
        if label in seen:
            continue
        seen.add(label)
        unique_tags.append((label, color))
    return unique_tags


def _format_opportunity_tags_text(opp: ScanOpportunity) -> str:
    """Tablo görünümü için etiket metnini birleştirir."""
    return " · ".join(label for label, _ in _get_opportunity_tags(opp))


def _render_opportunity_badges(opp: ScanOpportunity) -> None:
    """Coin adının yanına Streamlit badge etiketleri basar."""
    tags = _get_opportunity_tags(opp)
    st.markdown(f"#### {opp.symbol}")
    badge_cols = st.columns(min(len(tags), 4) or 1)
    for idx, (label, color) in enumerate(tags):
        with badge_cols[idx % len(badge_cols)]:
            if hasattr(st, "badge"):
                st.badge(label, color=color)
            else:
                st.markdown(
                    f"<span style='display:inline-block;padding:0.25rem 0.65rem;"
                    f"border-radius:999px;background:rgba(255,255,255,0.08);"
                    f"font-size:0.82rem;'>{escape(label)}</span>",
                    unsafe_allow_html=True,
                )


def build_opportunity_card_html(opp: ScanOpportunity) -> str:
    """LONG/SHORT fırsat kartı HTML'i oluşturur."""
    signal_label = _opp_signal_label(opp)
    is_long = signal_label == "LONG"
    rsi_text = f"{opp.rsi:.1f}" if opp.rsi is not None else "—"
    atr_text = f"{opp.atr:.4f}" if getattr(opp, "atr", None) else "—"
    tp1 = format_price(opp.take_profit_1 or 0) if opp.take_profit_1 else "—"
    tp2 = format_price(opp.take_profit_2 or 0) if opp.take_profit_2 else "—"
    sl = format_price(opp.stop_loss or 0) if opp.stop_loss else "—"
    tp1_eta = getattr(opp, "tp1_eta_text", None)
    tp2_eta = getattr(opp, "tp2_eta_text", None)
    tp1_eta_html = (
        f'<span class="eta-note">⏱ {escape(tp1_eta)}</span>' if tp1_eta else ""
    )
    tp2_eta_html = (
        f'<span class="eta-note">⏱ {escape(tp2_eta)}</span>' if tp2_eta else ""
    )
    profit_class = "green" if is_long else "red"
    sl_class = "red" if is_long else "green"
    accent = "#6ee7b7" if is_long else "#fca5a5"

    tags_html = " ".join(
        f'<span style="display:inline-block;margin:0.15rem 0.25rem;padding:0.2rem 0.55rem;'
        f'border-radius:999px;font-size:0.72rem;font-weight:600;background:rgba(255,255,255,0.08);'
        f'color:{accent};">{escape(label)}</span>'
        for label, _ in _get_opportunity_tags(opp)
    )

    return f"""
{opportunity_card_styles(is_long)}
<div class="opportunity-card">
    <div class="opp-header">
        <div class="opp-symbol">{escape(opp.symbol)}</div>
        <div class="opp-badge">{escape(signal_label)} · RSI {escape(rsi_text)}</div>
    </div>
    <div style="margin-bottom:0.85rem;">{tags_html}</div>
    <div class="opp-grid">
        <div class="opp-stat">
            <label>Güncel Fiyat</label>
            <span class="val">{escape(format_price(opp.current_price))}</span>
        </div>
        <div class="opp-stat">
            <label>Tavsiye Edilen Giriş</label>
            <span class="val {profit_class}">{escape(format_price(opp.entry_price))}</span>
        </div>
        <div class="opp-stat">
            <label>Kar Hedefi 1</label>
            <span class="val {profit_class}">{escape(tp1)}</span>
            {tp1_eta_html}
        </div>
        <div class="opp-stat">
            <label>Kar Hedefi 2</label>
            <span class="val {profit_class}">{escape(tp2)}</span>
            {tp2_eta_html}
        </div>
        <div class="opp-stat">
            <label>Zarar Durdur · ATR</label>
            <span class="val {sl_class}">{escape(sl)} · {escape(atr_text)}</span>
        </div>
    </div>
</div>
"""


def render_opportunity_card(opp: ScanOpportunity) -> None:
    """Tek fırsat kartını Streamlit bileşenleriyle renk kodlu basar."""
    signal_label = _opp_signal_label(opp)
    is_long = signal_label == "LONG"
    reason = getattr(opp, "reason", None) or opp.recommendation
    raw = getattr(opp, "raw_signal", None)
    is_early_long = bool(getattr(raw, "is_early_long", False)) if raw else False

    _render_opportunity_badges(opp)

    if is_early_long:
        st.warning(
            "⚠️ **Riskli / Erken LONG** — MACD teyidi henüz gelmemiş olabilir; "
            "dikkatli pozisyon boyutu önerilir.",
            icon="⚠️",
        )
    elif is_long:
        st.success("Pozisyon yönü: **Uzun (LONG)**", icon="✅")
    else:
        st.error("Pozisyon yönü: **Kısa (SHORT)**", icon="🔻")

    render_reason_block(reason)
    eta_badges: list[str] = []
    if getattr(opp, "tp1_eta_text", None):
        eta_badges.append(f":green-badge[KH1 · {opp.tp1_eta_text}]")
    if getattr(opp, "tp2_eta_text", None):
        eta_badges.append(f":blue-badge[KH2 · {opp.tp2_eta_text}]")
    if eta_badges:
        interval_label = getattr(opp, "interval", None) or "1h"
        st.markdown(" ".join(eta_badges))
        st.caption(
            f":material/schedule: Tahmini süre, ATR oynaklığı ve `{interval_label}` "
            f"zaman dilimindeki ortalama mum sayısına göredir."
        )

    render_html(build_opportunity_card_html(opp), height=340)

    render_add_to_watchlist_form(
        symbol=opp.symbol,
        default_entry=float(opp.entry_price or opp.current_price or 0),
        default_target=float(opp.take_profit_1 or opp.entry_price or 0),
        default_stop=float(opp.stop_loss) if opp.stop_loss else None,
        signal_label=signal_label,
        key_prefix=f"opp_{opp.symbol}",
    )


def _style_radar_dataframe(df: pd.DataFrame) -> pd.io.formats.style.Styler:
    """Özet tabloda LONG/SHORT sinyallerini renklendirir."""

    def _row_style(row: pd.Series) -> list[str]:
        sinyal = str(row.get("Yön", "")).upper()
        if sinyal == "LONG":
            bg = "background-color: rgba(16, 185, 129, 0.18); color: #6ee7b7; font-weight: 600"
        elif sinyal == "SHORT":
            bg = "background-color: rgba(239, 68, 68, 0.18); color: #fca5a5; font-weight: 600"
        else:
            bg = ""
        return [bg] * len(row)

    return df.style.apply(_row_style, axis=1)


def build_radar_export_dataframe(
    opportunities: list[ScanOpportunity],
) -> pd.DataFrame:
    """Aktif fırsatlar tablosunu Excel dışa aktarımı için hazırlar."""
    df = opportunities_to_dataframe(opportunities)
    df["etiketler"] = [_format_opportunity_tags_text(opp) for opp in opportunities]

    export_df = df[
        [
            "parite",
            "etiketler",
            "sinyal",
            "guncel_fiyat",
            "degisim_24s_yuzde",
            "tavsiye_giris",
            "kar_hedefi_1",
            "kar_hedefi_2",
            "kh1_tahmini_sure",
            "kh2_tahmini_sure",
            "zarar_durdur",
            "rsi",
            "atr",
            "gerekce",
            "taranma_zamani",
        ]
    ].copy()

    if "taranma_zamani" in export_df.columns:
        export_df["taranma_zamani"] = pd.to_datetime(
            export_df["taranma_zamani"], errors="coerce"
        ).dt.tz_localize(None)

    export_df = export_df.rename(
        columns={
            "parite": "Parite",
            "etiketler": "Etiketler",
            "sinyal": "Yön",
            "guncel_fiyat": "Güncel Fiyat",
            "degisim_24s_yuzde": "24s Değişim %",
            "tavsiye_giris": "Tavsiye Giriş",
            "kar_hedefi_1": "Kar Hedefi 1",
            "kar_hedefi_2": "Kar Hedefi 2",
            "kh1_tahmini_sure": "KH1 Tahmini Süre",
            "kh2_tahmini_sure": "KH2 Tahmini Süre",
            "zarar_durdur": "Zarar Durdur",
            "rsi": "RSI (14)",
            "atr": "ATR (14)",
            "gerekce": "Gerekçe",
            "taranma_zamani": "Taranma Zamanı",
        }
    )
    return export_df


def opportunities_to_excel_bytes(opportunities: list[ScanOpportunity]) -> bytes:
    """Fırsat listesini biçimlendirilmiş Excel dosyası olarak döndürür."""
    export_df = build_radar_export_dataframe(opportunities)
    buffer = BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="Aktif Fırsatlar")
        worksheet = writer.sheets["Aktif Fırsatlar"]

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="F9FAFB", bold=True)
        long_fill = PatternFill("solid", fgColor="D1FAE5")
        short_fill = PatternFill("solid", fgColor="FEE2E2")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        yon_col_idx = export_df.columns.get_loc("Yön") + 1
        gerekce_col_idx = export_df.columns.get_loc("Gerekçe") + 1

        for row_idx in range(2, worksheet.max_row + 1):
            yon_cell = worksheet.cell(row=row_idx, column=yon_col_idx)
            row_fill = None
            if str(yon_cell.value).upper() == "LONG":
                row_fill = long_fill
            elif str(yon_cell.value).upper() == "SHORT":
                row_fill = short_fill

            for col_idx in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if row_fill is not None:
                    cell.fill = row_fill
                if col_idx == gerekce_col_idx:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                else:
                    cell.alignment = Alignment(vertical="center")

        for col_idx, column_name in enumerate(export_df.columns, start=1):
            series = export_df[column_name].astype(str)
            max_len = max(series.map(len).max(), len(column_name)) + 2
            if column_name == "Gerekçe":
                width = 70
            elif column_name == "Etiketler":
                width = 28
            else:
                width = min(max_len, 22)
            worksheet.column_dimensions[get_column_letter(col_idx)].width = width

        worksheet.freeze_panes = "A2"

    buffer.seek(0)
    return buffer.getvalue()


def _render_watchlist_card_grid(
    symbols: list[str],
    snapshots: dict[str, CoinSnapshot],
) -> None:
    """İzleme listesi fiyat kartlarını çizer (fragment içinde kullanılır)."""
    if not snapshots:
        st.warning("Henüz gösterilecek veri yok. Lütfen birkaç saniye bekleyin.")
        return

    for row_start in range(0, len(symbols), 2):
        cols = st.columns(2)
        row_symbols = symbols[row_start : row_start + 2]

        for col_idx, symbol in enumerate(row_symbols):
            with cols[col_idx]:
                position = st.session_state.get("watchlist_positions", {}).get(symbol)
                if position:
                    entry_p = float(position.get("entry_price", 0))
                    target_p = float(position.get("target_price", 0))
                    stop_p = position.get("stop_loss")
                    stop_text = (
                        format_price(float(stop_p))
                        if stop_p is not None
                        else "—"
                    )
                    st.markdown(
                        f"<div style='background:rgba(96,165,250,0.08);border:1px solid "
                        f"rgba(96,165,250,0.25);border-radius:12px;padding:0.65rem 0.85rem;"
                        f"margin-bottom:0.65rem;font-size:0.82rem;color:#cbd5e1;'>"
                        f"📌 <strong>İzleme kaydı</strong> · Alış: "
                        f"{escape(format_price(entry_p))} · Hedef: "
                        f"{escape(format_price(target_p))} · SL: {escape(stop_text)}"
                        f"</div>",
                        unsafe_allow_html=True,
                    )
                snapshot = snapshots.get(symbol)
                if snapshot:
                    render_metric_card(snapshot)
                else:
                    render_html(
                        f"""
{card_styles()}
<div class="error-card">
    <strong>{escape(symbol)}</strong><br>
    Veri henüz yüklenmedi.
</div>
""",
                        height=120,
                    )


@st.fragment(run_every=REFRESH_INTERVAL_SEC, key="watchlist_live_cards")
def render_live_price_cards(
    symbols: list[str],
    interval: str,
    history_limit: int,
) -> None:
    """
    Yalnızca fiyat kartlarını 10 saniyede bir yeniler.

    Tam sayfa rerun yapmaz; yan menü, başlık ve radar sekmesi yerinde kalır.
    """
    refresh_count = int(st.session_state.get("watchlist_refresh_count", 0)) + 1
    st.session_state.watchlist_refresh_count = refresh_count

    try:
        snapshots = refresh_all_snapshots(
            symbols=symbols,
            interval=interval,
            history_limit=history_limit,
        )
        st.session_state.snapshots = snapshots
        st.session_state.last_refresh_at = utc_now()
    except Exception:
        st.error("Veri yenileme döngüsü başarısız oldu. Önceki veriler gösteriliyor.")
        snapshots = st.session_state.get("snapshots", {})

    has_stale = any(s.is_stale for s in snapshots.values()) if snapshots else False
    render_status_bar(refresh_count=refresh_count, has_stale=has_stale)
    _render_watchlist_card_grid(symbols, snapshots)


def render_watchlist_tab(
    symbols: list[str],
    interval: str,
    history_limit: int,
) -> None:
    """İzleme listesi sekmesini çizer."""
    if not symbols:
        st.info("Başlamak için yan menüden en az bir parite ekleyin.")
        return

    if st.session_state.get("scan_in_progress", False):
        st.info("Piyasa taraması devam ediyor. İzleme listesi tarama bitince güncellenecek.")
        snapshots = st.session_state.get("snapshots", {})
        if snapshots:
            render_status_bar(
                refresh_count=int(st.session_state.get("watchlist_refresh_count", 0)),
                has_stale=False,
            )
            _render_watchlist_card_grid(symbols, snapshots)
        return

    render_live_price_cards(
        symbols=symbols,
        interval=interval,
        history_limit=history_limit,
    )


def run_market_scan(interval: str, history_limit: int) -> tuple[list[ScanOpportunity], ScanStats]:
    """Piyasa taramasını ilerleme çubuğu ile çalıştırır."""
    config = ScannerConfig(
        pool_category_size=20,
        interval=interval,
        history_limit=history_limit,
        request_delay_sec=0.1,
        use_smart_pool=True,
    )

    bos_istatistik = ScanStats()

    ilerleme_alani = st.empty()
    durum_alani = st.empty()

    with ilerleme_alani.container():
        ilerleme_cubugu = st.progress(
            0,
            text="Akıllı Havuz hazırlanıyor (Günün en hareketli coinleri)...",
        )

    durum_alani.info(
        "🔍 **Akıllı Havuz Taranıyor** (Günün en hareketli coinleri)... "
        "Lütfen sayfayı kapatmayın.",
        icon="🚀",
    )

    def on_progress(current: int, total: int, symbol: str) -> None:
        yuzde = min(int((current / total) * 100), 100) if total else 0
        ilerleme_cubugu.progress(
            yuzde / 100,
            text=(
                f"Akıllı Havuz Taranıyor — {symbol} "
                f"({current}/{total}) · %{yuzde}"
            ),
        )
        durum_alani.markdown(
            f"🔥 **Günün en hareketli coinleri taranıyor…** `{symbol}` · "
            f"**{current} / {total}** parite analiz edildi"
        )

    try:
        opportunities, stats = scan_market(
            config=config,
            progress_callback=on_progress,
        )
        opportunities = _filter_radar_opportunities(opportunities)
    except Exception:
        durum_alani.empty()
        ilerleme_cubugu.progress(0, text="Tarama durduruldu.")
        st.warning(
            "Tarama sırasında bir hata oluştu, ancak devam ediliyor. "
            "Mümkünse birkaç saniye sonra tekrar deneyin.",
            icon="⚠️",
        )
        return st.session_state.get("scan_opportunities", []), bos_istatistik

    ilerleme_cubugu.progress(1.0, text="✅ Akıllı Havuz taraması tamamlandı!")

    if stats.atlanan > 0:
        st.warning(
            f"Tarama sırasında {stats.atlanan} paritede sorun oluştu, "
            f"ancak tarama tamamlandı. Sonuçlar aşağıda listeleniyor.",
            icon="⚠️",
        )

    if opportunities:
        long_n = stats.long_sinyali
        short_n = stats.short_sinyali
        durum_alani.success(
            f"Tarama tamamlandı — **{len(opportunities)}** aktif sinyal: "
            f"🟢 **{long_n} LONG**, 🔴 **{short_n} SHORT**. "
            f"(İncelenen: {stats.incelenen}, atlanan: {stats.atlanan})"
        )
    elif stats.incelenen == 0:
        durum_alani.warning(
            stats.mesaj
            or (
                "Tarama sonucu boş. Havuz oluşturulamadı veya taranan "
                "paritelerin hiçbiri için veri alınamadı."
            ),
            icon="⚠️",
        )
    else:
        durum_alani.info(
            stats.mesaj
            or (
                f"Tarama tamamlandı — LONG veya SHORT sinyali bulunamadı. "
                f"(İncelenen: {stats.incelenen}, atlanan: {stats.atlanan})"
            ),
            icon="ℹ️",
        )

    return opportunities, stats


def _apply_live_levels(
    opp: ScanOpportunity,
    live_price: float,
    atr_value: Optional[float],
    signal_label: str,
    interval: str,
) -> dict[str, Any]:
    """Anlık fiyat ve ATR ile kar hedefi / zarar durdur / ETA üretir."""
    levels = None
    if live_price > 0 and atr_value and atr_value > 0 and signal_label in ("LONG", "SHORT"):
        try:
            levels = calculate_exit_levels(live_price, signal_label, atr_value)
        except Exception:
            levels = getattr(opp, "exit_levels", None)

    tp1 = levels.take_profit_1 if levels else opp.take_profit_1
    tp2 = levels.take_profit_2 if levels else opp.take_profit_2
    sl = levels.stop_loss if levels else opp.stop_loss
    tp1_eta = (
        estimate_target_eta(live_price, tp1, atr_value, interval)
        if tp1 and atr_value
        else None
    )
    tp2_eta = (
        estimate_target_eta(live_price, tp2, atr_value, interval)
        if tp2 and atr_value
        else None
    )
    return {
        "exit_levels": levels or opp.exit_levels,
        "take_profit_1": tp1,
        "take_profit_2": tp2,
        "stop_loss": sl,
        "tp1_eta_text": tp1_eta.text if tp1_eta else opp.tp1_eta_text,
        "tp2_eta_text": tp2_eta.text if tp2_eta else opp.tp2_eta_text,
    }


def refresh_radar_opportunities(
    opportunities: list[ScanOpportunity],
    interval: str,
    history_limit: int,
) -> list[ScanOpportunity]:
    """
    Radar kartlarındaki fiyat, hedef ve durumu Binance anlık verisiyle yeniler.

    Önce tek istekle son fiyatlar alınır; ardından her coin için strateji
    ve ATR tabanlı kar hedefleri yeniden hesaplanır.
    """
    if not opportunities:
        return []

    symbols = [opp.symbol for opp in opportunities]
    try:
        tickers = fetch_live_tickers(symbols)
    except Exception:
        tickers = {}

    refreshed: list[ScanOpportunity] = []
    for opp in opportunities:
        ticker = tickers.get(opp.symbol.upper(), {})
        live_price = float(ticker.get("last_price") or opp.current_price or 0)
        change_pct = ticker.get("price_change_pct")
        quote_volume = ticker.get("quote_volume")
        atr_value = opp.atr
        signal_label = _opp_signal_label(opp)
        signal_type = opp.signal
        raw_signal = opp.raw_signal
        rsi = opp.rsi
        reason = opp.reason

        try:
            df = fetch_ohlcv(
                symbol=opp.symbol,
                interval=interval or getattr(opp, "interval", "1h"),
                limit=history_limit,
            )
            if df is not None and not df.empty:
                candle_close = float(df["Close"].iloc[-1])
                if live_price <= 0 and candle_close > 0:
                    live_price = candle_close
                signal_result = evaluate_signal(
                    df,
                    price_change_pct_24h=(
                        float(change_pct)
                        if change_pct is not None
                        else opp.price_change_pct_24h
                    ),
                )
                raw_signal = signal_result
                if signal_result.rsi is not None:
                    rsi = signal_result.rsi
                if signal_result.atr is not None:
                    atr_value = signal_result.atr
                if signal_result.signal_label in ("LONG", "SHORT"):
                    signal_label = signal_result.signal_label
                    signal_type = signal_result.signal
                if signal_result.message:
                    reason = signal_result.message
                time.sleep(0.05)
        except Exception:
            pass

        level_fields = _apply_live_levels(
            opp,
            live_price=live_price,
            atr_value=atr_value,
            signal_label=signal_label,
            interval=interval or getattr(opp, "interval", "1h"),
        )
        refreshed.append(
            replace(
                opp,
                current_price=live_price,
                rsi=rsi,
                atr=atr_value,
                volume_24h_usdt=(
                    float(quote_volume)
                    if quote_volume
                    else opp.volume_24h_usdt
                ),
                price_change_pct_24h=(
                    float(change_pct)
                    if change_pct is not None
                    else opp.price_change_pct_24h
                ),
                signal=signal_type,
                signal_label=signal_label,
                reason=reason,
                raw_signal=raw_signal,
                interval=interval or getattr(opp, "interval", "1h"),
                **level_fields,
            )
        )

    return refreshed


def _render_radar_results(
    opportunities: list[ScanOpportunity],
    scan_stats: Optional[ScanStats],
) -> None:
    """Radar özet tablosu ve tavsiye kartlarını çizer."""
    long_opps = [o for o in opportunities if _opp_signal_label(o) == "LONG"]
    short_opps = [o for o in opportunities if _opp_signal_label(o) == "SHORT"]

    pool_size = scan_stats.toplam if scan_stats else len(opportunities)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Taranan Havuz", pool_size)
    m2.metric("🟢 LONG", len(long_opps))
    m3.metric("🔴 SHORT", len(short_opps))
    if scan_stats:
        m4.metric(
            "Havuz Kaynağı",
            f"T{scan_stats.trend_sayisi}/D{scan_stats.loser_sayisi}/H{scan_stats.hacim_patlamasi_sayisi}",
            help="Trend / Düşen / Hacim patlaması kategorilerindeki parite sayıları",
        )
    else:
        m4.metric("Toplam Sinyal", len(opportunities))

    table_title_col, table_export_col = st.columns([3, 1])
    with table_title_col:
        st.markdown(f"#### Aktif Fırsatlar ({len(opportunities)} sinyal)")
    with table_export_col:
        export_filename = (
            f"sinyalci_firsatlar_{utc_now().strftime('%Y%m%d_%H%M')}.xlsx"
        )
        st.download_button(
            label="📥 Excel İndir",
            data=opportunities_to_excel_bytes(opportunities),
            file_name=export_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
            key="download_radar_excel",
        )

    display_df = build_radar_export_dataframe(opportunities).drop(
        columns=["Taranma Zamanı"]
    )
    st.dataframe(
        _style_radar_dataframe(display_df),
        hide_index=True,
    )

    st.markdown("---")

    if long_opps:
        st.markdown("#### 🔥 Uzun (LONG) Fırsatları")
        for row_start in range(0, len(long_opps), 2):
            cols = st.columns(2)
            for col_idx, opp in enumerate(long_opps[row_start : row_start + 2]):
                with cols[col_idx]:
                    render_opportunity_card(opp)

    if short_opps:
        if long_opps:
            st.markdown("---")
        st.markdown("#### 📉 Kısa (SHORT) Fırsatları")
        for row_start in range(0, len(short_opps), 2):
            cols = st.columns(2)
            for col_idx, opp in enumerate(short_opps[row_start : row_start + 2]):
                with cols[col_idx]:
                    render_opportunity_card(opp)


@st.fragment(run_every=REFRESH_INTERVAL_SEC, key="radar_live_cards")
def render_live_radar_cards(interval: str, history_limit: int) -> None:
    """
    Radar tavsiye kartlarını 10 saniyede bir Binance anlık verisiyle yeniler.

    Tam sayfa rerun yapmaz; tarama butonu ve geçmiş yerinde kalır.
    """
    opportunities = _filter_radar_opportunities(
        st.session_state.get("scan_opportunities", [])
    )
    scan_stats: Optional[ScanStats] = st.session_state.get("scan_stats")

    if not opportunities:
        bos_mesaj = (scan_stats.mesaj if scan_stats else "") or (
            "Şu an piyasada LONG veya SHORT sinyali bulunamadı — "
            "BEKLE modunda kalmanız önerilir."
        )
        if scan_stats and scan_stats.incelenen == 0:
            st.warning(bos_mesaj, icon="⚠️")
        else:
            st.info(bos_mesaj, icon="ℹ️")
        return

    refresh_count = int(st.session_state.get("radar_refresh_count", 0)) + 1
    st.session_state.radar_refresh_count = refresh_count

    try:
        opportunities = refresh_radar_opportunities(
            opportunities,
            interval=interval,
            history_limit=history_limit,
        )
        st.session_state.scan_opportunities = opportunities
        st.session_state.radar_last_refresh_at = utc_now()
    except Exception:
        st.warning(
            "Radar fiyatları yenilenemedi. Önceki tarama verileri gösteriliyor.",
            icon="⚠️",
        )

    last_live = st.session_state.get("radar_last_refresh_at")
    live_text = format_datetime_tr(last_live) if last_live else "—"
    st.caption(
        f":material/sync: Canlı radar · her {REFRESH_INTERVAL_SEC} sn · "
        f"döngü {refresh_count} · son fiyat: {live_text}"
    )
    _render_radar_results(opportunities, scan_stats)


def render_radar_tab(interval: str, history_limit: int) -> None:
    """Piyasa radarı ve tavsiyeler sekmesini çizer."""
    st.markdown("### Piyasa Radarı — Akıllı Coin Havuzu")
    st.caption(
        "Trend olanlar (Top 20 yükselen), düşen bıçaklar (Top 20 düşen) ve "
        "hacim patlaması (Top 20 yüksek hacim / düşük fiyat şişmesi) birleştirilerek "
        "24s hacim ve fiyat değişimine göre o an hareketli coinler dinamik olarak "
        "dahil edilir (en az 60 benzersiz USDT paritesi). Stablecoin/fiat pariteler kara listededir."
    )

    col_btn, col_info = st.columns([1, 2])
    with col_btn:
        scan_clicked = st.button(
            "Piyasayı Şimdi Tara (Akıllı Havuz ~60 Coin)",
            type="primary",
            use_container_width=True,
        )
    with col_info:
        last_scan = st.session_state.get("last_scan_at")
        scan_stats: Optional[ScanStats] = st.session_state.get("scan_stats")
        opp_count = len(st.session_state.get("scan_opportunities", []))
        if last_scan:
            stats_text = ""
            if scan_stats:
                stats_text = (
                    f" · Havuz: {scan_stats.toplam} coin"
                    f" · 🟢 {scan_stats.long_sinyali} LONG · "
                    f"🔴 {scan_stats.short_sinyali} SHORT"
                )
            st.caption(
                f"Son tarama: {format_datetime_tr(last_scan)} · "
                f"Toplam fırsat: {opp_count}{stats_text}"
            )
        else:
            st.caption("Henüz tarama yapılmadı. Butona basarak başlatın.")

    if scan_clicked:
        st.session_state.scan_in_progress = True
        try:
            opportunities, stats = run_market_scan(interval, history_limit)
            st.session_state.scan_opportunities = opportunities
            st.session_state.scan_stats = stats
            st.session_state.last_scan_at = utc_now()
            _update_recommendation_history(opportunities)
        finally:
            st.session_state.scan_in_progress = False

    render_recommendation_history()

    scan_stats: Optional[ScanStats] = st.session_state.get("scan_stats")

    if not st.session_state.get("last_scan_at"):
        st.info(
            "Akıllı coin havuzunu taramak için yukarıdaki butona tıklayın. "
            "Tarama birkaç dakika sürebilir.",
            icon="🔍",
        )
        return

    if scan_stats:
        st.caption(
            f"Son havuz: **{scan_stats.toplam}** benzersiz coin · "
            f"Dinamik: {getattr(scan_stats, 'dinamik_sayisi', 0)} · "
            f"Trend: {scan_stats.trend_sayisi} · "
            f"Düşen: {scan_stats.loser_sayisi} · "
            f"Hacim patlaması: {scan_stats.hacim_patlamasi_sayisi}"
        )

    if st.session_state.get("scan_in_progress", False):
        st.info("Tarama sürüyor. Canlı fiyat yenilemesi tarama bitince başlar.")
        return

    st.markdown("#### 🎯 Güncel Tarama Sonuçları")
    render_live_radar_cards(interval=interval, history_limit=history_limit)


# ---------------------------------------------------------------------------
# Ana uygulama
# ---------------------------------------------------------------------------


def render_sidebar() -> tuple[list[str], str, int]:
    """Yan menüyü çizer ve kullanıcı ayarlarını döndürür."""
    st.sidebar.markdown("## ⚙️ Ayarlar")

    if "symbols" not in st.session_state:
        st.session_state.symbols = ["BTCUSDT", "ETHUSDT"]

    coin_input = st.sidebar.text_input(
        "Coin Listesi (virgülle ayırın)",
        value=", ".join(st.session_state.symbols),
        placeholder="BTCUSDT, ETHUSDT, SOLUSDT",
        help="Takip etmek istediğiniz pariteleri virgülle yazın.",
    )

    zaman_dilimi_etiketleri = {
        "1m": "1 dakika",
        "5m": "5 dakika",
        "15m": "15 dakika",
        "1h": "1 saat",
        "4h": "4 saat",
        "1d": "1 gün",
    }
    interval = st.sidebar.selectbox(
        "Zaman Dilimi",
        options=list(zaman_dilimi_etiketleri.keys()),
        index=3,
        format_func=lambda x: zaman_dilimi_etiketleri[x],
    )

    history_limit = st.sidebar.slider(
        "Geçmiş Mum Sayısı",
        min_value=100,
        max_value=1000,
        value=500,
        step=50,
    )

    if st.sidebar.button("Listeyi Güncelle", use_container_width=True):
        parsed = parse_symbols(coin_input)
        if parsed:
            st.session_state.symbols = parsed
            st.session_state.pop("snapshots", None)
            st.rerun()
        else:
            st.sidebar.warning("En az bir geçerli coin girin.")

    st.sidebar.markdown("---")
    if st.sidebar.button("🔄 Şimdi Yenile", use_container_width=True):
        st.session_state.pop("snapshots", None)
        st.rerun()

    st.sidebar.markdown(
        dedent(
            f"""
            <p style="color:#64748b; font-size:0.78rem; margin-top:1rem;">
            Veriler Binance borsasından çekilir.<br>
            Fiyat kartları: <strong>{REFRESH_INTERVAL_SEC} saniyede bir</strong>
            (sayfa yenilenmez)
            </p>
            """
        ).strip(),
        unsafe_allow_html=True,
    )

    return st.session_state.symbols, interval, history_limit


def render_status_bar(refresh_count: int, has_stale: bool) -> None:
    """Son yenileme zamanını ve otomatik yenileme durumunu gösterir."""
    last_refresh = st.session_state.get("last_refresh_at")
    time_text = (
        format_datetime_tr(last_refresh)
        if last_refresh
        else "—"
    )
    dot_class = "status-dot stale" if has_stale else "status-dot"

    render_html(
        f"""
<div class="status-bar">
    <span class="{dot_class}"></span>
    <span>Kartlar 10 sn'de bir yenilenir (sayfa sabit) ·
    Döngü no: {refresh_count} · Son güncelleme: {escape(time_text)}</span>
</div>
"""
    )


def main() -> None:
    """Dashboard ana giriş noktası."""
    st.set_page_config(
        page_title="Sinyalci Paneli",
        page_icon="📊",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    hide_streamlit_chrome()

    inject_custom_css()
    ensure_recommendation_history_loaded()
    ensure_watchlist_loaded()

    if "symbols" not in st.session_state:
        st.session_state.symbols = ["BTCUSDT", "ETHUSDT"]
        # Diskteki izleme pozisyonları varsa sembolleri birleştir
        for sym in st.session_state.watchlist_positions:
            st.session_state.symbols = _add_symbol_to_watchlist(
                st.session_state.symbols,
                sym,
            )

    render_html(
        """
<div class="dashboard-header">
    <h1>Sinyalci</h1>
    <p>Kripto sinyal takip paneli — canlı fiyat, strateji ve risk seviyeleri</p>
</div>
"""
    )

    symbols, interval, history_limit = render_sidebar()

    tab_watchlist, tab_radar = st.tabs(
        ["👀 İzleme Listem", "🚀 Piyasa Radarı & Tavsiyeler"],
        on_change="rerun",
        key="main_tabs",
    )

    watch_open = tab_watchlist.open
    radar_open = tab_radar.open
    if watch_open is None and radar_open is None:
        watch_open = True

    if watch_open:
        with tab_watchlist:
            render_watchlist_tab(
                symbols=symbols,
                interval=interval,
                history_limit=history_limit,
            )

    if radar_open:
        with tab_radar:
            render_radar_tab(interval=interval, history_limit=history_limit)


if __name__ == "__main__":
    main()
